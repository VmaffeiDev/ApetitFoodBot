import logging
import os
import re
import sqlite3
import tempfile
import unicodedata
from datetime import UTC, date, datetime, timedelta
from html import escape
from pathlib import Path

import openpyxl
from dotenv import load_dotenv
from telegram import InlineKeyboardButton, InlineKeyboardMarkup, Update
from telegram.constants import ParseMode
from telegram.ext import (
    Application,
    CallbackQueryHandler,
    CommandHandler,
    ContextTypes,
    MessageHandler,
    filters,
)

load_dotenv()
logging.basicConfig(format="%(asctime)s - %(levelname)s - %(message)s", level=logging.INFO)
logger = logging.getLogger(__name__)

DB_PATH = Path(os.getenv("APETIT_DB_PATH", "apetit.db"))
DEFAULT_USER_NAME = os.getenv("APETIT_USER_NAME", "Colaborador")

REGISTRATION_STEP = "registration_step"
ADMIN_TELEGRAM_IDS = {
    int(value.strip())
    for value in os.getenv("ADMIN_TELEGRAM_IDS", "").split(",")
    if value.strip().isdigit()
}

RESTRICTIONS = {
    "restriction_none": "Sem restricoes",
    "restriction_gluten": "Sem gluten",
    "restriction_lactose": "Sem lactose",
    "restriction_vegetarian": "Vegetariana",
    "restriction_seafood": "Sem frutos do mar",
}

GOALS = {
    "goal_weight_loss": "Perder peso",
    "goal_muscle_gain": "Ganhar massa",
    "goal_maintenance": "Manter equilibrio",
    "goal_health": "Alimentacao mais saudavel",
    "goal_practical": "Praticidade na rotina",
}

WEEKDAY_LABELS = ["Segunda", "Terca", "Quarta", "Quinta", "Sexta", "Sabado", "Domingo"]


def normalize(text: str) -> str:
    return unicodedata.normalize("NFKD", text or "").encode("ascii", "ignore").decode().lower().strip()


def now_iso() -> str:
    return datetime.now(UTC).isoformat(timespec="seconds")


def format_date_br(iso_date: str) -> str:
    parsed = date.fromisoformat(iso_date)
    return f"{WEEKDAY_LABELS[parsed.weekday()]} ({parsed.strftime('%d/%m')})"


def db() -> sqlite3.Connection:
    conn = sqlite3.connect(DB_PATH)
    conn.row_factory = sqlite3.Row
    return conn


def init_db() -> None:
    with db() as conn:
        conn.executescript(
            """
            CREATE TABLE IF NOT EXISTS employees (
                telegram_id INTEGER PRIMARY KEY,
                chat_id INTEGER NOT NULL,
                name TEXT NOT NULL,
                goal TEXT NOT NULL DEFAULT 'Nao informado',
                restriction TEXT NOT NULL DEFAULT 'Sem restricoes',
                consent_accepted INTEGER NOT NULL DEFAULT 0,
                consented_at TEXT,
                created_at TEXT NOT NULL,
                updated_at TEXT NOT NULL
            );
            CREATE TABLE IF NOT EXISTS dishes (
                dish_key TEXT PRIMARY KEY,
                dish_name TEXT NOT NULL,
                base_category TEXT NOT NULL,
                ingredients TEXT NOT NULL DEFAULT '',
                allergens TEXT NOT NULL DEFAULT '',
                tags TEXT NOT NULL DEFAULT '',
                created_at TEXT NOT NULL,
                updated_at TEXT NOT NULL
            );
            CREATE TABLE IF NOT EXISTS daily_menu (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                menu_date TEXT NOT NULL,
                slot TEXT NOT NULL,
                base_category TEXT NOT NULL,
                dish_key TEXT NOT NULL,
                dish_name TEXT NOT NULL,
                created_at TEXT NOT NULL,
                UNIQUE (menu_date, slot)
            );
            CREATE TABLE IF NOT EXISTS selections (
                telegram_id INTEGER NOT NULL,
                menu_date TEXT NOT NULL,
                base_category TEXT NOT NULL,
                dish_key TEXT NOT NULL,
                dish_name TEXT NOT NULL,
                created_at TEXT NOT NULL,
                updated_at TEXT NOT NULL,
                PRIMARY KEY (telegram_id, menu_date, base_category)
            );
            """
        )


# ---------------------------------------------------------------------------
# Dishes
# ---------------------------------------------------------------------------

def parse_menu_cell(raw) -> tuple[str, str] | None:
    text = str(raw).strip() if raw is not None else ""
    if not text:
        return None
    parts = [part.strip() for part in text.split(" - ")]
    if len(parts) >= 3:
        code = parts[1]
        tail = parts[2:]
        if len(tail) > 1 and re.match(r"^[\d.,]+$", tail[-1]):
            tail = tail[:-1]
        name = " - ".join(tail).strip() or code
        return code, name
    return normalize(text)[:40] or "prato", text


def base_category_from_slot(slot: str) -> str:
    return re.sub(r"\s+\d+$", "", slot).strip().upper()


def upsert_dish_name(conn: sqlite3.Connection, code: str, name: str, base_category: str, timestamp: str) -> bool:
    existing = conn.execute("SELECT dish_key FROM dishes WHERE dish_key = ?", (code,)).fetchone()
    if existing:
        conn.execute("UPDATE dishes SET dish_name = ?, updated_at = ? WHERE dish_key = ?", (name, timestamp, code))
        return False
    conn.execute(
        """
        INSERT INTO dishes (dish_key, dish_name, base_category, created_at, updated_at)
        VALUES (?, ?, ?, ?, ?)
        """,
        (code, name, base_category, timestamp, timestamp),
    )
    return True


def import_menu_workbook(file_path: str, year: int, month: int) -> dict:
    workbook = openpyxl.load_workbook(file_path, data_only=True)
    sheet = workbook.worksheets[0]
    rows = list(sheet.iter_rows(values_only=True))
    if not rows:
        return {"days": 0, "new_dishes": [], "dates": []}
    header = [str(cell).strip() if cell else "" for cell in rows[0]]
    slot_columns = [(index, name) for index, name in enumerate(header) if index != 0 and name]
    timestamp = now_iso()
    imported_days = 0
    new_dishes: list[tuple[str, str]] = []
    imported_dates: list[str] = []
    with db() as conn:
        for row in rows[1:]:
            day_raw = row[0] if row else None
            if day_raw in (None, ""):
                continue
            try:
                day = int(str(day_raw).strip())
                menu_date = date(year, month, day).isoformat()
            except (ValueError, TypeError):
                continue
            row_has_item = False
            for col_index, slot in slot_columns:
                if col_index >= len(row):
                    continue
                parsed = parse_menu_cell(row[col_index])
                if not parsed:
                    continue
                code, name = parsed
                base_category = base_category_from_slot(slot)
                if upsert_dish_name(conn, code, name, base_category, timestamp):
                    new_dishes.append((code, name))
                conn.execute(
                    """
                    INSERT INTO daily_menu (menu_date, slot, base_category, dish_key, dish_name, created_at)
                    VALUES (?, ?, ?, ?, ?, ?)
                    ON CONFLICT(menu_date, slot) DO UPDATE SET
                        base_category = excluded.base_category,
                        dish_key = excluded.dish_key,
                        dish_name = excluded.dish_name,
                        created_at = excluded.created_at
                    """,
                    (menu_date, slot, base_category, code, name, timestamp),
                )
                row_has_item = True
            if row_has_item:
                imported_days += 1
                imported_dates.append(menu_date)
    return {"days": imported_days, "new_dishes": new_dishes, "dates": imported_dates}


def load_dish(dish_key: str) -> sqlite3.Row | None:
    with db() as conn:
        return conn.execute("SELECT * FROM dishes WHERE dish_key = ?", (dish_key,)).fetchone()


def update_dish_metadata(dish_key: str, ingredients: str, allergens: str, tags: str) -> None:
    with db() as conn:
        conn.execute(
            "UPDATE dishes SET ingredients = ?, allergens = ?, tags = ?, updated_at = ? WHERE dish_key = ?",
            (ingredients, allergens, tags, now_iso(), dish_key),
        )


def pending_dishes() -> list[sqlite3.Row]:
    with db() as conn:
        return conn.execute(
            "SELECT * FROM dishes WHERE ingredients = '' AND allergens = '' ORDER BY dish_name"
        ).fetchall()


def item_search_text(item: sqlite3.Row) -> str:
    return normalize(
        " ".join(
            [
                item["dish_name"] or "",
                item["ingredients"] or "",
                item["allergens"] or "",
                item["tags"] or "",
            ]
        )
    )


def restriction_conflict(restriction: str, item: sqlite3.Row) -> str | None:
    text = item_search_text(item)
    for safe_phrase in ("sem gluten", "sem lactose", "sem carne", "vegetariano", "vegetariana", "vegano", "vegana"):
        text = text.replace(safe_phrase, "")
    restriction_key = normalize(restriction)
    checks = {
        "sem gluten": (("gluten", "trigo", "farinha"), "contem gluten"),
        "sem lactose": (("lactose", "leite", "queijo", "creme", "manteiga"), "contem lactose ou derivados de leite"),
        "vegetariana": (("carne", "frango", "peixe", "bovina", "suina", "porco", "camarao"), "nao esta marcado como vegetariano"),
        "sem frutos do mar": (("peixe", "camarao", "frutos do mar", "marisco"), "contem peixe ou frutos do mar"),
    }
    for key, (terms, reason) in checks.items():
        if key in restriction_key and any(term in text for term in terms):
            return reason
    return None


def goal_score(goal: str, item: sqlite3.Row) -> int:
    text = item_search_text(item)
    goal_key = normalize(goal)
    keywords = {
        "perder peso": ("leve", "salada", "legumes", "sopa", "perda de peso", "grelhado", "fruta"),
        "ganhar massa": ("proteico", "frango", "carne", "ovo", "integral", "ganho de massa"),
        "manter equilibrio": ("equilibrado", "integral", "caseiro", "legumes", "natural"),
        "alimentacao mais saudavel": ("saudavel", "vegano", "integral", "legumes", "natural", "sem gluten", "fruta"),
        "praticidade na rotina": ("pratico", "kit", "rapido"),
    }
    for key, terms in keywords.items():
        if key in goal_key:
            return sum(1 for term in terms if term in text)
    return 0


def goal_reason(goal: str) -> str:
    goal_key = normalize(goal)
    if "perder peso" in goal_key:
        return "ajuda no seu objetivo de perder peso com uma opcao mais leve"
    if "ganhar massa" in goal_key:
        return "tem perfil mais proteico para apoiar ganho de massa"
    if "manter equilibrio" in goal_key:
        return "combina com uma rotina equilibrada"
    if "alimentacao mais saudavel" in goal_key:
        return "prioriza uma escolha mais natural e saudavel"
    if "praticidade na rotina" in goal_key:
        return "funciona bem para uma rotina mais pratica"
    return "combina com seu cadastro"


def best_dish_for_goal(dishes: list[sqlite3.Row], restriction: str, goal: str) -> sqlite3.Row | None:
    compatible = [dish for dish in dishes if not restriction_conflict(restriction, dish)]
    if not compatible:
        return None
    return max(compatible, key=lambda dish: goal_score(goal, dish))


# ---------------------------------------------------------------------------
# Weekly menu / selections
# ---------------------------------------------------------------------------

def current_week_dates() -> list[str]:
    today = date.today()
    monday = today - timedelta(days=today.weekday())
    return [(monday + timedelta(days=offset)).isoformat() for offset in range(7)]


def week_menu_dates(week_dates: list[str]) -> list[str]:
    placeholders = ",".join("?" for _ in week_dates)
    with db() as conn:
        rows = conn.execute(
            f"SELECT DISTINCT menu_date FROM daily_menu WHERE menu_date IN ({placeholders}) ORDER BY menu_date",
            week_dates,
        ).fetchall()
    return [row["menu_date"] for row in rows]


def day_menu_by_category(menu_date: str) -> dict[str, list[sqlite3.Row]]:
    with db() as conn:
        rows = conn.execute(
            """
            SELECT dm.base_category, dm.slot, d.*
            FROM daily_menu dm
            JOIN dishes d ON d.dish_key = dm.dish_key
            WHERE dm.menu_date = ?
            ORDER BY dm.base_category, dm.slot
            """,
            (menu_date,),
        ).fetchall()
    grouped: dict[str, list[sqlite3.Row]] = {}
    for row in rows:
        grouped.setdefault(row["base_category"], []).append(row)
    return grouped


def employee_selections(telegram_id: int, menu_date: str) -> dict[str, str]:
    with db() as conn:
        rows = conn.execute(
            "SELECT base_category, dish_key FROM selections WHERE telegram_id = ? AND menu_date = ?",
            (telegram_id, menu_date),
        ).fetchall()
    return {row["base_category"]: row["dish_key"] for row in rows}


def save_selection(telegram_id: int, menu_date: str, base_category: str, dish_key: str, dish_name: str) -> None:
    timestamp = now_iso()
    with db() as conn:
        conn.execute(
            """
            INSERT INTO selections (telegram_id, menu_date, base_category, dish_key, dish_name, created_at, updated_at)
            VALUES (?, ?, ?, ?, ?, ?, ?)
            ON CONFLICT(telegram_id, menu_date, base_category) DO UPDATE SET
                dish_key = excluded.dish_key,
                dish_name = excluded.dish_name,
                updated_at = excluded.updated_at
            """,
            (telegram_id, menu_date, base_category, dish_key, dish_name, timestamp, timestamp),
        )


def week_selections(telegram_id: int, week_dates: list[str]) -> list[sqlite3.Row]:
    placeholders = ",".join("?" for _ in week_dates)
    with db() as conn:
        return conn.execute(
            f"""
            SELECT menu_date, base_category, dish_name
            FROM selections
            WHERE telegram_id = ? AND menu_date IN ({placeholders})
            ORDER BY menu_date, base_category
            """,
            (telegram_id, *week_dates),
        ).fetchall()


# ---------------------------------------------------------------------------
# Employees
# ---------------------------------------------------------------------------

def save_employee(telegram_id: int, chat_id: int, name: str, goal: str, restriction: str, consent_accepted: bool, consented_at: str) -> None:
    timestamp = now_iso()
    with db() as conn:
        conn.execute(
            """
            INSERT INTO employees (telegram_id, chat_id, name, goal, restriction, consent_accepted, consented_at, created_at, updated_at)
            VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?)
            ON CONFLICT(telegram_id) DO UPDATE SET
                chat_id = excluded.chat_id,
                name = excluded.name,
                goal = excluded.goal,
                restriction = excluded.restriction,
                consent_accepted = excluded.consent_accepted,
                consented_at = excluded.consented_at,
                updated_at = excluded.updated_at
            """,
            (telegram_id, chat_id, name, goal, restriction, int(consent_accepted), consented_at, timestamp, timestamp),
        )


def load_employee(telegram_id: int) -> sqlite3.Row | None:
    with db() as conn:
        return conn.execute("SELECT * FROM employees WHERE telegram_id = ?", (telegram_id,)).fetchone()


def delete_employee_data(telegram_id: int) -> None:
    with db() as conn:
        conn.execute("DELETE FROM selections WHERE telegram_id = ?", (telegram_id,))
        conn.execute("DELETE FROM employees WHERE telegram_id = ?", (telegram_id,))


def list_active_employees() -> list[sqlite3.Row]:
    with db() as conn:
        return conn.execute("SELECT * FROM employees WHERE consent_accepted = 1").fetchall()


def restriction_conflicts_for_dates(restriction: str, menu_dates: list[str]) -> list[tuple[str, str, str]]:
    conflicts = []
    if not restriction:
        return conflicts
    for menu_date in menu_dates:
        for dishes in day_menu_by_category(menu_date).values():
            for dish in dishes:
                reason = restriction_conflict(restriction, dish)
                if reason:
                    conflicts.append((menu_date, dish["dish_name"], reason))
    return conflicts


def pending_dish_names_for_dates(menu_dates: list[str]) -> list[str]:
    seen: set[str] = set()
    names: list[str] = []
    for menu_date in menu_dates:
        for dishes in day_menu_by_category(menu_date).values():
            for dish in dishes:
                if not dish["ingredients"] and not dish["allergens"] and dish["dish_key"] not in seen:
                    seen.add(dish["dish_key"])
                    names.append(dish["dish_name"])
    return names


def admin_report_data() -> dict:
    with db() as conn:
        totals = conn.execute(
            """
            SELECT
                (SELECT COUNT(*) FROM employees) AS employees,
                (SELECT COUNT(*) FROM selections) AS selections,
                (SELECT COUNT(*) FROM dishes) AS dishes,
                (SELECT COUNT(*) FROM dishes WHERE ingredients = '' AND allergens = '') AS pending
            """
        ).fetchone()
        goals = conn.execute(
            "SELECT goal, COUNT(*) AS total FROM employees GROUP BY goal ORDER BY total DESC, goal"
        ).fetchall()
        top = conn.execute(
            """
            SELECT dish_name, COUNT(*) AS total
            FROM selections
            GROUP BY dish_key, dish_name
            ORDER BY total DESC, dish_name
            LIMIT 5
            """
        ).fetchall()
    return {"totals": totals, "goals": goals, "top": top}


# ---------------------------------------------------------------------------
# Telegram helpers
# ---------------------------------------------------------------------------

def tg_id(update: Update) -> int | None:
    return update.effective_user.id if update.effective_user else None


def chat_id(update: Update) -> int | None:
    return update.effective_chat.id if update.effective_chat else None


def profile(context: ContextTypes.DEFAULT_TYPE) -> dict:
    return context.user_data.setdefault("profile", {})


def hydrate_profile(update: Update, context: ContextTypes.DEFAULT_TYPE) -> None:
    user_id = tg_id(update)
    if not user_id or profile(context).get("registered"):
        return
    employee = load_employee(user_id)
    if employee and employee["consent_accepted"]:
        profile(context).update(
            {
                "telegram_id": employee["telegram_id"],
                "name": employee["name"],
                "goal": employee["goal"],
                "restriction": employee["restriction"],
                "registered": True,
            }
        )


def registered(context: ContextTypes.DEFAULT_TYPE) -> bool:
    return bool(profile(context).get("registered"))


def user_name(context: ContextTypes.DEFAULT_TYPE) -> str:
    return profile(context).get("name") or DEFAULT_USER_NAME


def render(text: str, context: ContextTypes.DEFAULT_TYPE) -> str:
    return text.format(name=escape(user_name(context)))


def keyboard(rows: list[list[tuple[str, str]]]) -> InlineKeyboardMarkup:
    return InlineKeyboardMarkup([[InlineKeyboardButton(label, callback_data=data) for label, data in row] for row in rows])


async def send_text(update: Update, text: str, context: ContextTypes.DEFAULT_TYPE, buttons=None, edit: bool = False) -> None:
    markup = keyboard(buttons) if buttons else None
    text = render(text, context)
    if edit and update.callback_query:
        await update.callback_query.edit_message_text(text=text, parse_mode=ParseMode.HTML, reply_markup=markup)
    elif update.effective_message:
        await update.effective_message.reply_text(text=text, parse_mode=ParseMode.HTML, reply_markup=markup)


async def send_payload(update: Update, payload: dict, context: ContextTypes.DEFAULT_TYPE, edit: bool = False) -> None:
    await send_text(update, payload["text"], context, payload.get("buttons"), edit)


def main_buttons() -> list[list[tuple[str, str]]]:
    return [
        [("\U0001f4c5 Cardapio da semana", "week_menu")],
        [("\U0001f464 Meu perfil", "profile")],
    ]


def restriction_buttons() -> list[list[tuple[str, str]]]:
    return [
        [("✅ Nenhuma restricao", "restriction_none")],
        [("\U0001f33e Sem gluten", "restriction_gluten"), ("\U0001f95b Sem lactose", "restriction_lactose")],
        [("\U0001f969 Vegetariana", "restriction_vegetarian"), ("\U0001f420 Sem frutos do mar", "restriction_seafood")],
    ]


def goal_buttons() -> list[list[tuple[str, str]]]:
    return [
        [("\U0001f343 Perder peso", "goal_weight_loss"), ("\U0001f4aa Ganhar massa", "goal_muscle_gain")],
        [("⚖️ Manter equilibrio", "goal_maintenance")],
        [("\U0001f957 Alimentacao saudavel", "goal_health"), ("⏱️ Praticidade", "goal_practical")],
    ]


def consent_buttons() -> list[list[tuple[str, str]]]:
    return [[("✅ Aceito", "consent_yes"), ("❌ Nao aceito", "consent_no")]]


# ---------------------------------------------------------------------------
# Payloads
# ---------------------------------------------------------------------------

def week_menu_payload(restriction: str) -> dict:
    week_dates = current_week_dates()
    dates_with_menu = week_menu_dates(week_dates)
    if not dates_with_menu:
        return {
            "text": "\U0001f614 Ainda nao recebemos o cardapio desta semana. Assim que a empresa enviar, aviso por aqui.",
            "buttons": [],
        }
    lines = ["\U0001f4c5 <b>Cardapio da semana:</b>"]
    buttons = []
    for menu_date in dates_with_menu:
        grouped = day_menu_by_category(menu_date)
        categories = []
        for base_category, dishes in grouped.items():
            names = ", ".join(escape(dish["dish_name"]) for dish in dishes)
            conflict_names = [dish["dish_name"] for dish in dishes if restriction_conflict(restriction, dish)]
            warn = " ⚠️" if restriction and conflict_names else ""
            categories.append(f"{escape(base_category.title())}: {names}{warn}")
        lines.append(f"\n<b>{escape(format_date_br(menu_date))}</b>\n" + "\n".join(categories))
        buttons.append([(f"Ver {format_date_br(menu_date)}", f"day:{menu_date}")])
    return {"text": "\n".join(lines), "buttons": buttons}


def day_menu_payload(menu_date: str, context: ContextTypes.DEFAULT_TYPE) -> dict:
    data = profile(context)
    restriction = data.get("restriction", "")
    goal = data.get("goal", "")
    telegram_id = data.get("telegram_id")
    grouped = day_menu_by_category(menu_date)
    if not grouped:
        return {"text": "\U0001f614 Nao encontrei cardapio para esse dia.", "buttons": [[("\U0001f4c5 Cardapio da semana", "week_menu")]]}
    selected = employee_selections(telegram_id, menu_date) if telegram_id else {}
    lines = [f"\U0001f4c5 <b>Cardapio de {escape(format_date_br(menu_date))}:</b>"]
    buttons = []
    for base_category, dishes in grouped.items():
        lines.append(f"\n<b>{escape(base_category.title())}</b>")
        best = best_dish_for_goal(dishes, restriction, goal) if goal else None
        for dish in dishes:
            conflict = restriction_conflict(restriction, dish) if restriction else None
            marker = ""
            if conflict:
                marker = f" ⚠️ {escape(conflict)}"
            elif best and dish["dish_key"] == best["dish_key"]:
                marker = " ⭐ recomendado para seu objetivo"
            picked = selected.get(base_category) == dish["dish_key"]
            check = " ✅" if picked else ""
            lines.append(f"- {escape(dish['dish_name'])}{check}{marker}")
            if not conflict:
                label = f"Escolher {dish['dish_name'][:26]}"
                buttons.append([(label, f"pick:{menu_date}:{base_category}:{dish['dish_key']}")])
    buttons.append([("\U0001f4c5 Cardapio da semana", "week_menu")])
    return {"text": "\n".join(lines), "buttons": buttons}


def profile_payload(context: ContextTypes.DEFAULT_TYPE) -> dict:
    data = profile(context)
    return {
        "text": (
            "\U0001f464 <b>Seu perfil:</b>\n\n"
            f"<b>Nome:</b> {escape(data.get('name', DEFAULT_USER_NAME))}\n"
            f"<b>Objetivo:</b> {escape(data.get('goal', 'Nao informado'))}\n"
            f"<b>Restricao alimentar:</b> {escape(data.get('restriction', 'Nao informado'))}"
        ),
        "buttons": [[("✏️ Atualizar cadastro", "restart_registration"), ("✅ Esta correto", "thanks")]],
    }


# ---------------------------------------------------------------------------
# Registration
# ---------------------------------------------------------------------------

async def ask_name(update: Update, context: ContextTypes.DEFAULT_TYPE, edit: bool = False) -> None:
    context.user_data[REGISTRATION_STEP] = "name"
    await send_text(
        update,
        (
            "\U0001f957 <b>Bem-vindo ao Apetit!</b>\n\n"
            "Vou te ajudar a acompanhar o cardapio da semana da empresa, recomendar pratos de acordo com seu objetivo "
            "e avisar quando algo tiver ingrediente que voce nao pode comer.\n\n"
            "Antes, preciso de um cadastro rapido. Qual e o seu nome?"
        ),
        context,
        edit=edit,
    )


async def ask_goal(update: Update, context: ContextTypes.DEFAULT_TYPE, edit: bool = False) -> None:
    context.user_data[REGISTRATION_STEP] = "goal"
    await send_text(update, "\U0001f3af Qual e o seu objetivo principal com a alimentacao agora?", context, goal_buttons(), edit=edit)


async def ask_restriction(update: Update, context: ContextTypes.DEFAULT_TYPE, edit: bool = False) -> None:
    context.user_data[REGISTRATION_STEP] = "restriction"
    await send_text(update, "\U0001f33f Voce tem alguma restricao ou alergia alimentar?", context, restriction_buttons(), edit=edit)


async def ask_consent(update: Update, context: ContextTypes.DEFAULT_TYPE, edit: bool = False) -> None:
    context.user_data[REGISTRATION_STEP] = "consent"
    await send_text(
        update,
        (
            "\U0001f512 <b>Antes de concluir:</b>\n\n"
            "Vou guardar seu nome, objetivo e restricao alimentar para poder te recomendar pratos e avisar sobre alergenos "
            "no cardapio da empresa. Voce pode consultar (/meus_dados) ou apagar (/excluir_dados) esses dados quando quiser.\n\n"
            "Voce aceita?"
        ),
        context,
        consent_buttons(),
        edit=edit,
    )


async def finish_registration(update: Update, context: ContextTypes.DEFAULT_TYPE, accepted: bool, edit: bool = False) -> None:
    data = profile(context)
    user_id = tg_id(update)
    current_chat_id = chat_id(update)
    if not accepted:
        context.user_data.pop(REGISTRATION_STEP, None)
        await send_text(
            update,
            "Sem problemas. Sem o aceite eu nao consigo guardar seu cadastro, entao nao vou liberar o cardapio agora. "
            "Quando quiser tentar de novo, envie /start.",
            context,
        )
        return
    consented_at = now_iso()
    if user_id:
        data["telegram_id"] = user_id
    data["registered"] = True
    if user_id and current_chat_id:
        save_employee(
            user_id,
            current_chat_id,
            data.get("name", DEFAULT_USER_NAME),
            data.get("goal", "Nao informado"),
            data.get("restriction", "Sem restricoes"),
            True,
            consented_at,
        )
    context.user_data.pop(REGISTRATION_STEP, None)
    await send_text(
        update,
        (
            "✅ <b>Cadastro concluido!</b>\n\n"
            "<b>Nome:</b> {name}\n"
            f"<b>Objetivo:</b> {escape(data.get('goal', 'Nao informado'))}\n"
            f"<b>Restricao alimentar:</b> {escape(data.get('restriction', 'Sem restricoes'))}\n\n"
            "Agora posso te mostrar o cardapio da semana \U0001f60a"
        ),
        context,
        main_buttons(),
        edit,
    )


async def start(update: Update, context: ContextTypes.DEFAULT_TYPE) -> None:
    hydrate_profile(update, context)
    if registered(context):
        await send_text(
            update,
            (
                "\U0001f957 <b>Ola, {name}!</b>\n\n"
                "Posso te mostrar o cardapio da semana, recomendar pratos para o seu objetivo e avisar sobre alergenos \U0001f33f"
            ),
            context,
            main_buttons(),
        )
        return
    await ask_name(update, context)


async def reset_registration(update: Update, context: ContextTypes.DEFAULT_TYPE) -> None:
    context.user_data.clear()
    await ask_name(update, context)


async def show_my_data(update: Update, context: ContextTypes.DEFAULT_TYPE) -> None:
    user_id = tg_id(update)
    employee = load_employee(user_id) if user_id else None
    if not employee:
        await update.effective_message.reply_text("Voce ainda nao tem cadastro. Envie /start para se cadastrar.")
        return
    await update.effective_message.reply_text(
        "\U0001f4cb Seus dados salvos:\n\n"
        f"Nome: {employee['name']}\n"
        f"Objetivo: {employee['goal']}\n"
        f"Restricao alimentar: {employee['restriction']}\n"
        f"Consentimento LGPD: aceito em {employee['consented_at']}\n\n"
        "Para apagar tudo, envie /excluir_dados."
    )


async def delete_my_data(update: Update, context: ContextTypes.DEFAULT_TYPE) -> None:
    user_id = tg_id(update)
    if not user_id:
        return
    delete_employee_data(user_id)
    context.user_data.clear()
    await update.effective_message.reply_text(
        "✅ Seus dados (cadastro e escolhas de cardapio) foram apagados. Se quiser voltar a usar o bot, envie /start."
    )


async def show_week_menu(update: Update, context: ContextTypes.DEFAULT_TYPE) -> None:
    hydrate_profile(update, context)
    if not registered(context):
        await ask_name(update, context)
        return
    await send_payload(update, week_menu_payload(profile(context).get("restriction", "")), context)


async def show_my_selections(update: Update, context: ContextTypes.DEFAULT_TYPE) -> None:
    hydrate_profile(update, context)
    user_id = tg_id(update)
    if not user_id or not registered(context):
        await ask_name(update, context)
        return
    rows = week_selections(user_id, current_week_dates())
    if not rows:
        await update.effective_message.reply_text("Voce ainda nao escolheu nenhum prato desta semana. Use /cardapio_semana.")
        return
    lines = ["\U0001f4cb Suas escolhas desta semana:"]
    last_date = None
    for row in rows:
        if row["menu_date"] != last_date:
            lines.append(f"\n{format_date_br(row['menu_date'])}")
            last_date = row["menu_date"]
        lines.append(f"- {row['base_category'].title()}: {row['dish_name']}")
    await update.effective_message.reply_text("\n".join(lines))


# ---------------------------------------------------------------------------
# Admin
# ---------------------------------------------------------------------------

def is_admin(user_id: int | None) -> bool:
    return bool(ADMIN_TELEGRAM_IDS) and user_id in ADMIN_TELEGRAM_IDS


async def broadcast_new_menu(context: ContextTypes.DEFAULT_TYPE, menu_dates: list[str]) -> int:
    if not menu_dates:
        return 0
    start_label = format_date_br(min(menu_dates))
    end_label = format_date_br(max(menu_dates))
    period = start_label if start_label == end_label else f"{start_label} a {end_label}"
    pending = pending_dish_names_for_dates(menu_dates)
    notified = 0
    for employee in list_active_employees():
        conflicts = restriction_conflicts_for_dates(employee["restriction"], menu_dates)
        lines = [
            "\U0001f4c5 <b>Cardapio novo disponivel!</b>",
            "",
            f"A empresa ja publicou o cardapio de {escape(period)}. De uma olhada no que vai ter e escolha o que voce vai comer.",
        ]
        if pending:
            lines.append("")
            lines.append(
                "ℹ️ Alguns pratos dessa semana ainda nao tem ingredientes/alergenicos cadastrados "
                f"({', '.join(escape(name) for name in pending[:5])}). Se tiver alergia, confirme com o refeitorio antes de comer."
            )
        if conflicts:
            lines.append("")
            lines.append(f"⚠️ <b>Atencao:</b> tem prato(s) essa semana que batem com sua restricao ({escape(employee['restriction'])}):")
            lines.extend(f"- {escape(format_date_br(menu_date))}: {escape(dish_name)} ({escape(reason)})" for menu_date, dish_name, reason in conflicts)
        try:
            await context.bot.send_message(
                chat_id=employee["chat_id"],
                text="\n".join(lines),
                parse_mode=ParseMode.HTML,
                reply_markup=keyboard([[("\U0001f4c5 Ver cardapio da semana", "week_menu")]]),
            )
            notified += 1
        except Exception:
            logger.exception("Falha ao notificar colaborador %s sobre novo cardapio", employee["telegram_id"])
    return notified


async def handle_menu_document(update: Update, context: ContextTypes.DEFAULT_TYPE) -> None:
    document = update.effective_message.document
    if not document or not (document.file_name or "").lower().endswith(".xlsx"):
        return
    if not is_admin(tg_id(update)):
        await update.effective_message.reply_text("\U0001f512 Apenas administradores podem importar o cardapio.")
        return
    caption = (update.effective_message.caption or "").strip()
    match = re.match(r"^(\d{1,2})/(\d{4})$", caption)
    today = date.today()
    month, year = (int(match.group(1)), int(match.group(2))) if match else (today.month, today.year)
    telegram_file = await document.get_file()
    with tempfile.NamedTemporaryFile(suffix=".xlsx", delete=False) as tmp:
        tmp_path = tmp.name
    try:
        await telegram_file.download_to_drive(tmp_path)
        result = import_menu_workbook(tmp_path, year, month)
    finally:
        Path(tmp_path).unlink(missing_ok=True)
    notified = await broadcast_new_menu(context, result["dates"])
    lines = [f"✅ Cardapio importado para {month:02d}/{year}: {result['days']} dia(s) com itens. Colaboradores notificados: {notified}."]
    if not match:
        lines.append("(Nenhum mes/ano na legenda da mensagem, usei o mes atual. Envie a legenda no formato MM/AAAA para escolher outro mes.)")
    if result["new_dishes"]:
        lines.append("")
        lines.append(f"⚠️ {len(result['new_dishes'])} prato(s) novo(s) sem ingredientes/alergenicos cadastrados:")
        lines.extend(f"- {code}: {name}" for code, name in result["new_dishes"][:15])
        lines.append("")
        lines.append("Complete cada um com:\n/prato_add <codigo> | ingredientes | alergenicos | tags")
    await update.effective_message.reply_text("\n".join(lines))


async def add_dish_metadata(update: Update, context: ContextTypes.DEFAULT_TYPE) -> None:
    if not is_admin(tg_id(update)):
        await update.effective_message.reply_text("\U0001f512 Apenas administradores podem cadastrar alergenicos.")
        return
    text = update.effective_message.text or ""
    raw = text.split(maxsplit=1)[1].strip() if len(text.split(maxsplit=1)) > 1 else ""
    parts = [part.strip() for part in raw.split("|")]
    if len(parts) < 4:
        await update.effective_message.reply_text(
            "Envie assim:\n\n/prato_add <codigo> | ingredientes | alergenicos | tags\n\n"
            "O codigo e o que aparece em /pratos_pendentes."
        )
        return
    code, ingredients, allergens, tags = parts[:4]
    dish = load_dish(code)
    if not dish:
        await update.effective_message.reply_text(f"Nao encontrei prato com codigo {code}. Importe o cardapio antes.")
        return
    update_dish_metadata(code, ingredients, allergens, tags)
    await update.effective_message.reply_text(f"✅ Prato atualizado: {dish['dish_name']} ({code}).")


async def list_pending_dishes(update: Update, context: ContextTypes.DEFAULT_TYPE) -> None:
    if not is_admin(tg_id(update)):
        await update.effective_message.reply_text("\U0001f512 Apenas administradores podem ver essa lista.")
        return
    pending = pending_dishes()
    if not pending:
        await update.effective_message.reply_text("✅ Nenhum prato pendente de cadastro de alergenicos.")
        return
    lines = ["⚠️ Pratos sem ingredientes/alergenicos cadastrados:"]
    lines.extend(f"- {dish['dish_key']}: {dish['dish_name']}" for dish in pending)
    await update.effective_message.reply_text("\n".join(lines))


async def show_admin_report(update: Update, context: ContextTypes.DEFAULT_TYPE) -> None:
    if not is_admin(tg_id(update)):
        await update.effective_message.reply_text("\U0001f512 Apenas administradores podem ver o relatorio.")
        return
    report = admin_report_data()
    totals = report["totals"]
    lines = [
        "\U0001f4ca Relatorio Apetit Bot",
        "",
        f"Colaboradores cadastrados: {totals['employees']}",
        f"Escolhas de cardapio registradas: {totals['selections']}",
        f"Pratos cadastrados: {totals['dishes']}",
        f"Pratos pendentes de alergenicos: {totals['pending']}",
        "",
        "Objetivos dos colaboradores:",
    ]
    if report["goals"]:
        lines.extend(f"- {row['goal']}: {row['total']}" for row in report["goals"])
    else:
        lines.append("- Ainda sem colaboradores cadastrados.")
    lines.extend(["", "Pratos mais escolhidos:"])
    if report["top"]:
        lines.extend(f"- {row['dish_name']}: {row['total']}" for row in report["top"])
    else:
        lines.append("- Ainda sem escolhas registradas.")
    await update.effective_message.reply_text("\n".join(lines))


# ---------------------------------------------------------------------------
# Message / callback handlers
# ---------------------------------------------------------------------------

async def handle_registration_message(update: Update, context: ContextTypes.DEFAULT_TYPE) -> bool:
    step = context.user_data.get(REGISTRATION_STEP)
    if not step or step in {"goal", "restriction", "consent"}:
        return step is not None
    text = (update.message.text or "").strip()
    if len(text) < 2:
        await send_text(update, "Me envie uma resposta um pouco mais completa, por favor \U0001f60a", context)
        return True
    data = profile(context)
    if step == "name":
        data["name"] = text
        await ask_goal(update, context)
    return True


async def handle_callback(update: Update, context: ContextTypes.DEFAULT_TYPE) -> None:
    hydrate_profile(update, context)
    query = update.callback_query
    await query.answer()
    data = query.data or ""

    if data in {"restart_registration", "update_profile"}:
        profile(context).pop("registered", None)
        await ask_name(update, context, edit=True)
        return
    if data in GOALS:
        profile(context)["goal"] = GOALS[data]
        await ask_restriction(update, context, edit=True)
        return
    if data in RESTRICTIONS:
        profile(context)["restriction"] = RESTRICTIONS[data]
        await ask_consent(update, context, edit=True)
        return
    if data in {"consent_yes", "consent_no"}:
        await finish_registration(update, context, accepted=(data == "consent_yes"), edit=True)
        return
    if data == "profile":
        if not registered(context):
            await ask_name(update, context, edit=True)
            return
        await send_payload(update, profile_payload(context), context, edit=True)
        return
    if data == "week_menu":
        if not registered(context):
            await ask_name(update, context, edit=True)
            return
        await send_payload(update, week_menu_payload(profile(context).get("restriction", "")), context, edit=True)
        return
    if data.startswith("day:"):
        if not registered(context):
            await ask_name(update, context, edit=True)
            return
        menu_date = data.removeprefix("day:")
        await send_payload(update, day_menu_payload(menu_date, context), context, edit=True)
        return
    if data.startswith("pick:"):
        if not registered(context):
            await ask_name(update, context, edit=True)
            return
        _, menu_date, base_category, dish_key = data.split(":", 3)
        user_id = tg_id(update)
        dish = load_dish(dish_key)
        if not dish:
            await send_payload(update, day_menu_payload(menu_date, context), context, edit=True)
            return
        conflict = restriction_conflict(profile(context).get("restriction", ""), dish)
        if conflict:
            await send_text(
                update,
                f"⚠️ <b>{escape(dish['dish_name'])}</b> {escape(conflict)} e pode nao combinar com sua restricao. Nao registrei essa escolha.",
                context,
                edit=True,
            )
            return
        if user_id:
            save_selection(user_id, menu_date, base_category, dish_key, dish["dish_name"])
        await send_payload(update, day_menu_payload(menu_date, context), context, edit=True)
        return
    await send_text(update, "Como posso ajudar? \U0001f60a", context, main_buttons(), edit=True)


async def handle_message(update: Update, context: ContextTypes.DEFAULT_TYPE) -> None:
    hydrate_profile(update, context)
    if await handle_registration_message(update, context):
        return
    if not registered(context):
        await ask_name(update, context)
        return
    text = normalize(update.message.text or "")
    if "cardapio" in text or "semana" in text:
        await send_payload(update, week_menu_payload(profile(context).get("restriction", "")), context)
    elif "perfil" in text or "cadastro" in text:
        await send_payload(update, profile_payload(context), context)
    elif "escolha" in text:
        await show_my_selections(update, context)
    else:
        await send_text(update, "Quer ver o cardapio da semana ou seu perfil? \U0001f60a", context, main_buttons())


def main() -> None:
    token = os.getenv("TELEGRAM_BOT_TOKEN")
    if not token:
        raise RuntimeError("Defina TELEGRAM_BOT_TOKEN no arquivo .env antes de iniciar o bot.")
    init_db()
    app = Application.builder().token(token).build()
    app.add_handler(CommandHandler("start", start))
    app.add_handler(CommandHandler("recadastrar", reset_registration))
    app.add_handler(CommandHandler("cardapio_semana", show_week_menu))
    app.add_handler(CommandHandler("minhas_escolhas", show_my_selections))
    app.add_handler(CommandHandler("meus_dados", show_my_data))
    app.add_handler(CommandHandler("excluir_dados", delete_my_data))
    app.add_handler(CommandHandler("prato_add", add_dish_metadata))
    app.add_handler(CommandHandler("pratos_pendentes", list_pending_dishes))
    app.add_handler(CommandHandler("relatorio", show_admin_report))
    app.add_handler(CallbackQueryHandler(handle_callback))
    app.add_handler(MessageHandler(filters.Document.FileExtension("xlsx"), handle_menu_document))
    app.add_handler(MessageHandler(filters.TEXT & ~filters.COMMAND, handle_message))
    logger.info("Apetit Bot iniciado.")
    app.run_polling(allowed_updates=Update.ALL_TYPES)


if __name__ == "__main__":
    main()
